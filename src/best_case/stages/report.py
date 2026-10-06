"""Stage - comparison report in the layout of the presentation "Best Case Demo".

Builds one table per slide (same column titles as the slide), an Excel workbook with
native charts and PNG charts.  Everything shown is configurable in ``settings.yaml``
(``report:`` section); nothing here changes the workflow result itself.

Definitions (slide 4 of the presentation)
-----------------------------------------
* Approach I  : per hole section the *minimum over wells* of the section's total duration.
* Approach II : per hole section the *sum over activities (code1)* of the minimum duration
                of that activity over wells.
* Approach Combination : per hole section the smaller of Approach I and II.
"""
from __future__ import annotations

import logging
import re
from pathlib import Path

import pandas as pd

from ..engine.expr import like_mask
from .base import Context

log = logging.getLogger("best_case")

SEC, WELL, TYPE, FORM, DAYS = "Hole Section", "Well", "Type", "Formation", "Duration (Days)"
DEFAULT_PARAMETERS = [
    "bittfa", "stringno", "depthend", "depthstart", "dttmend", "hookloadoffbottom", "hookloadpickup",
    "hookloadrotating", "hookloadslackoff", "liquidinjrate", "rpmmotor", "rpmstring", "sppdiff", "sppdrill",
    "tfo", "tmcirc", "tmdrill", "torquedrill", "torqueoffbtm", "wob", "des_DrillstringComp",
]


def hole_size(label) -> float:
    """``"12 1/4"`` -> 12.25 (used to order hole sections from large to small)."""
    m = re.match(r"\s*(\d+)(?:\s+(\d+)\s*/\s*(\d+))?", str(label))
    if not m:
        return -1.0
    size = float(m.group(1))
    if m.group(2):
        size += float(m.group(2)) / float(m.group(3))
    return size


def section_label(code4) -> str:
    return f"{code4}″"


def _prepare(sections: pd.DataFrame, p: dict) -> pd.DataFrame:
    s = sections.copy()
    s["days"] = s["duration"] / float(p.get("hours_per_day", 24))
    if p.get("wells"):
        s = s[s["wellname"].isin(p["wells"])]
    if p.get("hole_sections"):
        s = s[s["code4"].isin([str(x) for x in p["hole_sections"]])]
    return s


def _order(codes) -> list:
    return sorted(set(codes), key=hole_size, reverse=True)


def build_tables(t: dict[str, pd.DataFrame], cfg) -> dict[str, pd.DataFrame]:
    p = cfg.get("report", {}) or {}
    s = _prepare(t["sections"], p)
    if s.empty:
        raise ValueError("report: no section rows left after applying report.wells / report.hole_sections")
    order = _order(s["code4"])
    out: dict[str, pd.DataFrame] = {}

    # ---- slide 6: Approach I - duration by well & hole section
    per_well = s.groupby(["wellname", "code4", "category"], dropna=False)["days"].sum().reset_index()
    wide = per_well.pivot_table(index="code4", columns="wellname", values="days", aggfunc="sum").reindex(order)
    wide.insert(0, SEC, [section_label(c) for c in wide.index])
    out["Approach I by Well"] = wide.reset_index(drop=True).rename_axis(columns=None)

    # ---- approaches
    a1 = per_well.loc[per_well.groupby("code4")["days"].idxmin()].set_index("code4").reindex(order)
    per_act = s.dropna(subset=["code1"]).groupby(["wellname", "code4", "code1"])["days"].sum().reset_index()
    a2 = per_act.groupby(["code4", "code1"])["days"].min().groupby("code4").sum().reindex(order)
    comp = pd.DataFrame({
        SEC: [section_label(c) for c in order],
        "Approach I": a1["days"].to_numpy(),
        "Approach II": a2.to_numpy(),
    })
    comp["Approach Combination"] = comp[["Approach I", "Approach II"]].min(axis=1)
    out["Approach I vs II"] = comp[[SEC, "Approach I", "Approach II"]]
    total = pd.DataFrame([{SEC: "Total", **comp.drop(columns=SEC).sum().to_dict()}])
    out["Comparison"] = pd.concat([comp, total], ignore_index=True)
    out["All Approaches"] = comp
    out["Approach I winners"] = pd.DataFrame({SEC: comp[SEC], "Best Well (Approach I)": a1["wellname"].to_numpy(),
                                              TYPE: a1["category"].to_numpy(), DAYS: a1["days"].to_numpy()})

    # ---- slide 11: total duration by hole section in each formation
    f = s.groupby(["code4", "wellname", "category", "formname"], dropna=False)["days"].sum().reset_index()
    f["_o"] = f["code4"].map(hole_size)
    f = f.sort_values(["_o", "category", "days"], ascending=[False, True, True]).drop(columns="_o")
    out["Duration by Formation"] = pd.DataFrame({
        SEC: f["code4"].map(section_label), WELL: f["wellname"], TYPE: f["category"],
        FORM: f["formname"].astype(str).str.title(), DAYS: f["days"],
    })

    # ---- slide 12: minimum duration by hole section, formation and well type
    b = t["best_sections"].copy()
    b = b[b["code4"].isin(order)]
    b["_o"] = b["code4"].map(hole_size)
    b = b.sort_values(["_o", "category", "formname"], ascending=[False, True, True])
    out["Min Duration by Formation"] = pd.DataFrame({
        SEC: b["code4"].map(section_label), WELL: b["wellname"], TYPE: b["category"],
        FORM: b["formname"].astype(str).str.title(), DAYS: b["Sum(duration)"] / float(p.get("hours_per_day", 24)),
    })

    # ---- slide 13: drilling parameter extraction (I) - the DDR comments of the best case
    tl = t["best_time_log"]
    ddr = p.get("ddr", {}) or {}
    pat = ddr.get("comment_pattern", "*param*")
    keep = tl["code4"].isin(order)
    if pat:
        keep &= like_mask(tl["com"], pat, bool(ddr.get("case_sensitive", False)))
    d = tl[keep].sort_values(["wellname", "dttmend_date", "sysseq"])
    out["Drilling Parameters I"] = pd.DataFrame({
        SEC: d["code4"].map(section_label), WELL: d["wellname"], "Date": d["dttmend_date"].dt.date,
        FORM: d["formname"].astype(str).str.title(), "Duration (Hours)": d["duration"],
        "Drilling Parameters (DDRs)": d["com"],
    })

    # ---- slide 14: drilling parameter extraction (II) - one block per best-case parameter record
    res = t["best_case_result"]
    keys = tl[["wellname", "formname", "dttmend_date", "code4"]].drop_duplicates()
    r = res.merge(keys, on=["wellname", "formname", "dttmend_date"], how="inner")
    r = r[r["code4"].isin(order)]
    params = [c for c in (p.get("parameters") or DEFAULT_PARAMETERS) if c in r.columns]
    rows = []
    for rec, (_, row) in enumerate(r.iterrows(), start=1):
        for c in params:
            v = row[c]
            if isinstance(v, pd.Timestamp):
                v = v.date().isoformat()
            elif isinstance(v, float):
                v = None if pd.isna(v) else round(v, 2)
            rows.append({"Record": rec, SEC: section_label(row["code4"]), WELL: row["wellname"],
                         "Date": row["dttmend_date"].date(), FORM: str(row["formname"]).title(),
                         "Parameter": c, "Value": "-" if v is None or (isinstance(v, float) and pd.isna(v)) else v})
    out["Drilling Parameters II"] = pd.DataFrame(rows, columns=["Record", SEC, WELL, "Date", FORM, "Parameter", "Value"])
    return out


# ------------------------------------------------------------------------------------ Excel

# sheet name (<= 31 chars) -> slide it reproduces
SHEETS = {
    "Approach I by Well": "6 Approach I - Duration by Well",
    "Approach I vs II": "7 Approach I vs II",
    "Comparison": "8 Comparison",
    "All Approaches": "9 Duration - All Approaches",
    "Approach I winners": "9b Approach I winners",
    "Duration by Formation": "11 Duration by Formation",
    "Min Duration by Formation": "12 Min Duration by Formation",
    "Drilling Parameters I": "13 Drilling Parameters I",
    "Drilling Parameters II": "14 Drilling Parameters II",
}


def write_excel(tables: dict[str, pd.DataFrame], path: Path, p: dict) -> None:
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Font

    with pd.ExcelWriter(path, engine="openpyxl", date_format="yyyy-mm-dd") as w:
        for key, sheet in SHEETS.items():
            if key in tables:
                tables[key].to_excel(w, sheet_name=sheet, index=False)
        wb = w.book
        for ws in wb.worksheets:
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)
            ws.freeze_panes = "A2"

        def add_chart(sheet_key: str, title: str, y_title: str, anchor: str, last_row: int | None = None):
            ws = wb[SHEETS[sheet_key]]
            last = last_row or ws.max_row
            ch = BarChart()
            ch.type, ch.grouping = "col", "clustered"
            ch.title, ch.y_axis.title, ch.x_axis.title = title, y_title, "Hole Section"
            data = Reference(ws, min_col=2, max_col=ws.max_column, min_row=1, max_row=last)
            cats = Reference(ws, min_col=1, min_row=2, max_row=last)
            ch.add_data(data, titles_from_data=True)
            ch.set_categories(cats)
            ch.dataLabels = DataLabelList()
            ch.dataLabels.showVal = ws.max_column <= 5
            ch.y_axis.numFmt = "0.0"
            ch.height, ch.width = 9.5, 20
            ch.x_axis.delete = ch.y_axis.delete = False
            ws.add_chart(ch, anchor)

        if "Approach I by Well" in tables:
            add_chart("Approach I by Well", "Approach I - Duration by Well & Hole Section", "Days", "A12")
        if "Approach I vs II" in tables:
            add_chart("Approach I vs II", "Approach I vs. Approach II (Minimum Duration by Hole Section)", "Days", "F2")
        if "All Approaches" in tables:
            add_chart("All Approaches", "Duration by Hole Section - All Approaches", "Days", "F2")
        if "Comparison" in tables:
            ws = wb[SHEETS["Comparison"]]
            for row in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row):
                for c in row:
                    c.font = Font(bold=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.00"


# ------------------------------------------------------------------------------------ stage


def run(ctx: Context) -> None:
    cfg = ctx.cfg
    tables = build_tables(ctx.tables, cfg)
    out_dir = cfg.path("export.directory") / cfg.get("report.directory", "report")
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, df in tables.items():
        ctx.tables[f"report::{name}"] = df
    p = cfg.get("report", {}) or {}
    if p.get("excel", True):
        path = out_dir / p.get("excel_file", "best_case_report.xlsx")
        write_excel(tables, path, p)
        log.info("report workbook written: %s", path)
    if p.get("png", True):
        try:
            from . import report_charts

            files = report_charts.render_all(tables, ctx.tables, out_dir, p)
            log.info("report charts written: %s", ", ".join(f.name for f in files))
        except ImportError:
            log.warning("matplotlib is not installed - PNG charts skipped (pip install matplotlib)")
